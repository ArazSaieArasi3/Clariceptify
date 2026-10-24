"""Controlled import of the selected, hash-bound DOCX; no semantic adjudication."""
import argparse
from pathlib import Path
import json
import hashlib
import re
from collections import Counter
import yaml
from docx import Document
from jsonschema import Draft202012Validator

ROOT=Path(__file__).resolve().parents[1]


def generate(source):
    doc=Document(source)
    digest=hashlib.sha256(Path(source).read_bytes()).hexdigest()
    manifest=yaml.safe_load((ROOT/'traceability/migration/input-integrity.yaml').read_text())
    if digest != next(x['sha256'] for x in manifest['records'] if x['source_id']=='SRC-SPEC'):
        raise ValueError('Input digest differs from the frozen selected source; record a new source revision first.')
    def rows(index):return [[c.text.strip() for c in r.cells] for r in doc.tables[index].rows[1:]]
    cs,ds,members,rs=rows(1),rows(2),rows(3),rows(4)
    ledger=yaml.safe_load((ROOT/'traceability/migration/entity-ledger.yaml').read_text())
    key_by_label={r['selected_label']:r['proposed_record_id'] for r in ledger['selected_concepts']}
    domain_by_label={r[0]:f'CLR-D{i:04d}' for i,r in enumerate(ds,1)}
    membership={}
    for row in members:
        labels=row[2].split('; ')
        if len(labels)!=int(row[1]):raise ValueError('Input membership count mismatch: '+row[0])
        for label in labels:
            if label in membership:raise ValueError('Duplicate primary membership: '+label)
            membership[label]=domain_by_label[row[0]]
    if set(membership)!=set(key_by_label):raise ValueError('Membership and concept catalogue differ')
    def common(table,row):return {'lifecycle':'candidate','introduced_version':'input-r1-rev1',
                                  'source_locator':f'SRC-SPEC@sha256:{digest}#table-{table}/data-row-{row}'}
    domains=[dict(common(3,i),id=domain_by_label[r[0]],label_en=r[0],purpose=r[1],description=r[2],
                  layer='foundations' if r[0]=='Foundations' else 'subject') for i,r in enumerate(ds,1)]
    concepts=[dict(common(2,i),id=key_by_label[r[0]],label_en=r[0],ufo_type=r[1],ontouml_stereotype=r[2],
                   definition=r[3],is_abstract=None,primary_domain_id=membership[r[0]],aliases=[],source_reference_ids=[],
                   notes='Input reference display: '+r[4]+'. Stable source-support bindings await the source-register task. Historical correspondence remains unreviewed in the migration ledger.') for i,r in enumerate(cs,1)]
    associations=[];generalizations=[]
    for i,r in enumerate(rs,1):
        a,b=key_by_label[r[0]],key_by_label[r[8]]
        if r[3]=='specializes':
            generalizations.append(dict(common(5,i),id=f'CLR-G{len(generalizations)+1:04d}',child_id=a,parent_id=b,
                 set_completeness='unknown',set_disjointness='unknown',
                 notes='Input connector: '+r[5]+'. End stereotypes: '+r[1]+' / '+r[7]+'. No partition axiom inferred.'))
        else:
            associations.append(dict(common(5,i),id=f'CLR-R{len(associations)+1:04d}',label_en=r[3],source_id=a,target_id=b,
                 source_multiplicity=r[2],target_multiplicity=r[6],ontouml_stereotype=None if r[4]=='None' else r[4],
                 connector=r[5],notes='Input end stereotypes: '+r[1]+' / '+r[7]+'. Multiplicities are the labels at the respective diagram ends; no OWL restriction is inferred.'))
    start=next(i for i,p in enumerate(doc.paragraphs) if p.text.startswith('6 Semantic constraints'))
    end=next(i for i in range(start+1,len(doc.paragraphs)) if doc.paragraphs[i].text.startswith('7 Integration review'))
    texts=[p.text for p in doc.paragraphs[start+2:end] if p.text.strip()]
    constraints=[]
    for i,text in enumerate(texts,1):
        # Retain explicit lexical mentions only as candidates; no semantic closure inferred.
        mentions=[key_by_label[label] for label in key_by_label if re.search(r'(?<![A-Za-z])'+re.escape(label)+r'(?![A-Za-z])',text)]
        constraints.append(dict(id=f'CLR-SC{i:04d}',text=text,referenced_ids=mentions,enforcement='textual',
             execution_status='not_executed',lifecycle='candidate',introduced_version='input-r1-rev1',
             source_locator=f'SRC-SPEC@sha256:{digest}#section-6/constraint-{i}',
             notes='Referenced IDs are case-sensitive lexical mention candidates only; completeness and contextual interpretation require review. Text is preserved without translation or formalization.'))
    families={'domain':domains,'concept':concepts,'relation':associations,'generalization':generalizations,'constraint':constraints}
    expected={'domain':21,'concept':137,'relation':174,'generalization':122,'constraint':24}
    for kind,entries in families.items():
        if len(entries)!=expected[kind]:raise ValueError('Selected input count differs: '+kind)
        schema=json.loads((ROOT/f'model/schema/{kind}.schema.json').read_text())
        for record in entries:Draft202012Validator(schema).validate(record)
    assert all(record['definition']==row[3] and record['ontouml_stereotype']==row[2] for record,row in zip(concepts,cs))
    assert len(set(key_by_label.values()))==137
    assert len(set((r['source_id'],r['label_en'],r['target_id'],r['source_multiplicity'],r['target_multiplicity'],r['ontouml_stereotype']) for r in associations))==174
    assert len(set((r['child_id'],r['parent_id']) for r in generalizations))==122
    filenames={'domain':'domains','concept':'concepts','relation':'relations','generalization':'generalizations','constraint':'constraints'}
    for kind,entries in families.items():
        payload={'schema_version':'1.0','asset_id':'CLR-MODEL-'+kind.upper(),'model_revision':'input-r1-rev1',
                 'status':'imported_candidate','source_sha256':digest,'record_schema':f'schema/{kind}.schema.json','records':entries}
        (ROOT/f'model/{filenames[kind]}.yaml').write_text(yaml.safe_dump(payload,allow_unicode=True,sort_keys=False))
    result={'status':'PASS','scope':'controlled source extraction and schema parity; not semantic approval',
            'source_sha256':digest,'counts':{k:len(v) for k,v in families.items()},
            'concept_definitions_and_stereotypes':'preserved exactly from source table',
            'domain_names':'preserved exactly; 20 subject domains plus Foundations',
            'relationship_partition':'174 associations and 122 directed specializations; exact source-row locators retained',
            'constraint_text':'24 paragraphs preserved exactly; not executed',
            'unresolved_work':['source-reference stable bindings','diagram reconciliation','requirement/CQ alignment',
                               'semantic lineage decisions','full adequacy and satisfiability checks','independent review']}
    (ROOT/'evaluation/model-import-report.json').write_text(json.dumps(result,indent=2)+'\n')
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source_docx');args=parser.parse_args()
    print(json.dumps(generate(args.source_docx),indent=2))
