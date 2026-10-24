"""Extract Stage-1 citations and support from the frozen specification representation."""
from pathlib import Path
import re,json,hashlib,yaml,argparse
from docx import Document

ROOT=Path(__file__).resolve().parents[1]


def generate(path):
    digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    integrity=yaml.safe_load((ROOT/'traceability/migration/input-integrity.yaml').read_text())
    if digest!=next(x['sha256'] for x in integrity['records'] if x['source_id']=='SRC-SPEC'):raise ValueError('Selected source digest differs')
    doc=Document(path);refs=[]
    for i,p in enumerate(doc.paragraphs):
        match=re.match(r'^\[(\d+)\] (.+)',p.text)
        if match:
            n=int(match[1]);refs.append({'id':f'CLR-SRC{n:04d}','input_reference_number':n,'citation_ieee_input':p.text,
                'source_locator':f'SRC-SPEC@sha256:{digest}#paragraph-{i+1}',
                'metadata_verification':'not_reverified_in_import','full_text_review':'not_assessed'})
    assert [r['input_reference_number'] for r in refs]==list(range(1,54))
    register={'schema_version':'1.0','source_artifact':'SRC-SPEC','source_sha256':digest,
        'authority_note':'The Stage-1 source/support tables and bibliography in the frozen specification are the actual imported representation. Byte parity to the separately identified Stage-1 JSON ledger has not been established.',
        'identity_policy':'IDs initially allocated from this input numbering; future display renumbering must retain IDs and bibliographic identity. A material source replacement gets a new ID.',
        'records':refs}
    concepts=yaml.safe_load((ROOT/'model/concepts.yaml').read_text());ids={x['label_en']:x['id'] for x in concepts['records']}
    ref_ids={r['input_reference_number']:r['id'] for r in refs};links=[]
    refcells={r.cells[0].text.strip():r.cells[4].text.strip() for r in doc.tables[1].rows[1:]}
    for i,row in enumerate(doc.tables[7].rows[1:],1):
        name,kind,scope,locator=[c.text.strip() for c in row.cells]
        numbers=[int(n) for n in re.findall(r'\[(\d+)\]',refcells[name])]
        assert len(numbers)>=2 and all(n in ref_ids for n in numbers)
        links.append({'concept_id':ids[name],'input_label':name,'support_type':kind,
                      'source_ids':[ref_ids[n] for n in numbers],'input_reference_display':refcells[name],
                      'support_scope_and_limit':scope,'input_source_locations':locator,
                      'source_locator':f'SRC-SPEC@sha256:{digest}#table-8/data-row-{i}',
                      'passage_verification':'not_reverified_in_import','typing_approval':'not_assessed'})
    assert len(links)==137 and len({r['concept_id'] for r in links})==137
    from collections import Counter
    counts=dict(Counter(r['support_type'] for r in links));assert counts=={'B':83,'D':44,'P':10},counts
    for c in concepts['records']:c['source_reference_ids']=next(r['source_ids'] for r in links if r['concept_id']==c['id'])
    for c in concepts['records']:
        c['notes']='Input reference display: '+refcells[c['label_en']]+'. Stable bindings are in evidence/concept-support.yaml. Source support is qualified, not independent typing approval; historical correspondences remain unreviewed.'
    (ROOT/'evidence/source-register.yaml').write_text(yaml.safe_dump(register,allow_unicode=True,sort_keys=False))
    (ROOT/'evidence/concept-support.yaml').write_text(yaml.safe_dump({'schema_version':'1.0','source_sha256':digest,'qualification':{'D':'direct notion support in input screening','B':'adapted/theoretical basis','P':'specific passage support pending'},'records':links},allow_unicode=True,sort_keys=False))
    (ROOT/'evidence/bibliography.md').write_text('# Input bibliography\n\nImported from the frozen specification; metadata and full text have not been reverified in this import. Reference numbers are display values, not source identity.\n\n'+'\n\n'.join(r['citation_ieee_input'] for r in refs)+'\n')
    (ROOT/'model/concepts.yaml').write_text(yaml.safe_dump(concepts,allow_unicode=True,sort_keys=False))
    result={'status':'PASS','source_sha256':digest,'representation':'specification Stage-1 support table and bibliography',
            'references':53,'concept_support_records':137,'support_types':counts,
            'minimum_source_pointers_per_concept':min(len(r['source_ids']) for r in links),
            'limits':['no fresh primary-source or full-text verification','P entries remain pending','multiple pointers need not be independent','separate JSON ledger bytes not compared']}
    (ROOT/'evaluation/source-support-import.json').write_text(json.dumps(result,indent=2)+'\n');return result


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('source_docx');a=ap.parse_args();print(json.dumps(generate(a.source_docx),indent=2))
